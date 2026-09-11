"""What actually lands in the .xlsx file."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

import swdesigntables as sw


def test_title_cell(simple_table, reload_table):
    sheet = reload_table(simple_table)["Sheet1"]
    assert sheet["A1"].value == "Design Table for: BRK-MASTER"


def test_headers_start_at_b2_and_a2_stays_empty(simple_table, reload_table):
    """The classic off-by-one-column mistake, asserted head on.

    A2 carries the Family defined name and must stay empty; headers start at B2.
    """
    sheet = reload_table(simple_table)["Sheet1"]
    assert sheet["A2"].value is None
    assert sheet["B2"].value == "Length@Boss-Extrude1"
    assert sheet["C2"].value == "$VALUE@hole_qty@Equations"


def test_configuration_names_fill_column_a(simple_table, reload_table):
    sheet = reload_table(simple_table)["Sheet1"]
    assert sheet["A3"].value == "BRK-025"
    assert sheet["A4"].value == "BRK-040"


def test_values_land_under_their_own_header(simple_table, simple_columns, reload_table):
    """Look values up by header, never by column number."""
    sheet = reload_table(simple_table)["Sheet1"]
    headers = {cell.value: cell.column for cell in sheet[2] if cell.value}
    assert sheet.cell(row=3, column=headers["Length@Boss-Extrude1"]).value == 25.0
    assert sheet.cell(row=3, column=headers["$STATE@Draft2"]).value == "S"
    assert sheet.cell(row=4, column=headers["$STATE@Draft2"]).value == "U"
    assert sheet.cell(row=4, column=headers["$PRP@Material"]).value == "6061-T6"


def test_family_defined_name_is_workbook_scoped(simple_table, reload_table):
    """Without Family in A2, SOLIDWORKS cannot find the table at all."""
    workbook = reload_table(simple_table)
    assert "Family" in workbook.defined_names
    assert workbook.defined_names["Family"].attr_text == "Sheet1!$A$2"


def test_sheet_name_with_spaces_is_quoted_in_the_reference():
    table = sw.DesignTable("M", [sw.prop("Description")], sheet_name="Design Data")
    table.add_configuration("A", {sw.prop("Description"): "x"})
    buffer = io.BytesIO()
    table.save(buffer)
    buffer.seek(0)
    workbook = load_workbook(buffer)
    assert workbook.defined_names["Family"].attr_text == "'Design Data'!$A$2"


def test_swx_sheet_is_never_authored(simple_table, reload_table):
    """SOLIDWORKS creates _SWX itself when it embeds the table."""
    assert "_SWX" not in reload_table(simple_table).sheetnames


def test_expression_is_written_as_text_not_a_formula(reload_table):
    """A formula cell has no cached result, so SOLIDWORKS reads nothing."""
    shrink = sw.global_variable("shrink")
    table = sw.DesignTable("M", [shrink], ignore=("expression-in-global-variable",))
    table.add_configuration("A", {shrink: sw.Expression("rate * Length")})
    sheet = reload_table(table)["Sheet1"]
    assert sheet["B3"].value == "=rate * Length"
    assert sheet["B3"].data_type == "s"


def test_expression_normalizes_a_leading_equals():
    assert sw.Expression("=W/2").render() == "=W/2"
    assert sw.Expression("W/2").render() == "=W/2"


def test_numeric_looking_strings_stay_text(reload_table):
    """A part number of '0012' must not come back as the number 12."""
    part = sw.prop("PartNo")
    table = sw.DesignTable("M", [part])
    table.add_configuration("A", {part: "0012"})
    sheet = reload_table(table)["Sheet1"]
    assert sheet["B3"].value == "0012"
    assert isinstance(sheet["B3"].value, str)


def test_missing_value_defaults_to_blank(simple_columns, reload_table):
    columns = simple_columns
    table = sw.DesignTable("M", list(columns.values()))
    table.add_configuration("A", {columns["length"]: 10.0})
    sheet = reload_table(table)["Sheet1"]
    assert sheet["B3"].value == 10.0
    assert sheet["C3"].value is None


def test_missing_value_error_policy(simple_columns):
    columns = simple_columns
    table = sw.DesignTable("M", list(columns.values()), missing=sw.MissingValue.ERROR)
    table.add_configuration("A", {columns["length"]: 10.0})
    report = table.validate()
    assert "missing-value" in {issue.code for issue in report.errors}


def test_missing_value_fill_policy(simple_columns, reload_table):
    columns = simple_columns
    table = sw.DesignTable(
        "M", list(columns.values()), missing=sw.MissingValue.FILL, fill_value=0
    )
    table.add_configuration("A", {columns["length"]: 10.0})
    sheet = reload_table(table)["Sheet1"]
    assert sheet["C3"].value == 0


def test_numeric_state_format(reload_table):
    """Working tables in the wild use 1/0; the library must not rewrite them."""
    draft = sw.state("Draft2")
    table = sw.DesignTable("M", [draft], state_format=sw.StateFormat.NUMERIC)
    table.add_configuration("A", {draft: sw.SUPPRESSED})
    table.add_configuration("B", {draft: sw.UNSUPPRESSED})
    sheet = reload_table(table)["Sheet1"]
    assert sheet["B3"].value == 1
    assert sheet["B4"].value == 0


def test_round_floats(reload_table):
    length = sw.dimension("Length", "Boss-Extrude1")
    table = sw.DesignTable("M", [length], round_floats=3)
    table.add_configuration("A", {length: 316.66666666666663})
    sheet = reload_table(table)["Sheet1"]
    assert sheet["B3"].value == 316.667


def test_save_accepts_path_str_and_stream(tmp_path, simple_table):
    as_path = simple_table.save(tmp_path / "a.xlsx")
    assert as_path.exists()
    assert simple_table.save(str(tmp_path / "b.xlsx")).exists()
    buffer = io.BytesIO()
    assert simple_table.save(buffer) is None
    assert buffer.getvalue()


def test_save_uses_output_path_when_called_bare(tmp_path, simple_columns):
    columns = simple_columns
    destination = tmp_path / "nested" / "out.xlsx"
    table = sw.DesignTable("M", list(columns.values()), output_path=destination)
    table.add_configuration("A", {columns["length"]: 1.0})
    assert table.save() == destination
    assert destination.exists()


def test_save_without_target_or_output_path_is_refused(simple_table):
    with pytest.raises(ValueError, match="No target given"):
        simple_table.save()


def test_to_bytes_matches_a_saved_file(tmp_path, simple_table):
    payload = simple_table.to_bytes()
    written = (tmp_path / "x.xlsx")
    simple_table.save(written)
    assert payload == written.read_bytes()


def test_unsupported_value_type_names_the_problem():
    anything = sw.prop("Whatever")
    table = sw.DesignTable("M", [anything])
    table.add_configuration("A", {anything: object()})
    with pytest.raises(TypeError, match="Unsupported cell value"):
        table.to_bytes()


def test_decimal_is_accepted(reload_table):
    from decimal import Decimal

    length = sw.dimension("Length", "Boss-Extrude1")
    table = sw.DesignTable("M", [length])
    table.add_configuration("A", {length: Decimal("12.5")})
    assert reload_table(table)["Sheet1"]["B3"].value == 12.5
