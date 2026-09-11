"""Reusable bases: TableTemplate and blank_table."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

import swdesigntables as sw

LENGTH = sw.dimension("Length", "Boss-Extrude1")
WIDTH = sw.dimension("Width", "Boss-Extrude1")


def bracket_template(tmp_path):
    return sw.TableTemplate(
        model_name="BRK-MASTER",
        columns=(LENGTH, WIDTH, sw.description()),
        output_dir=tmp_path,
        file_name="brk_master_dt.xlsx",
        config_name=lambda width, length: f"BRK-{width:03d}-{length:03d}",
        round_floats=3,
    )


def test_new_table_inherits_model_columns_and_path(tmp_path):
    template = bracket_template(tmp_path)
    table = template.new_table()
    assert table.model_name == "BRK-MASTER"
    assert table.headers() == ("Length@Boss-Extrude1", "Width@Boss-Extrude1", "$DESCRIPTION")
    assert table.output_path == tmp_path / "brk_master_dt.xlsx"
    assert table.round_floats == 3


def test_save_with_no_argument_uses_the_template_path(tmp_path):
    template = bracket_template(tmp_path)
    table = template.new_table()
    table.add_configuration("BRK-025-040", {LENGTH: 40.0, WIDTH: 25.0})
    written = table.save()
    assert written == tmp_path / "brk_master_dt.xlsx"
    assert written.exists()


def test_name_for_uses_the_naming_rule(tmp_path):
    template = bracket_template(tmp_path)
    assert template.name_for(width=25, length=40) == "BRK-025-040"


def test_name_for_without_a_rule_says_so():
    template = sw.TableTemplate(model_name="M")
    with pytest.raises(ValueError, match="no config_name rule"):
        template.name_for()


def test_replace_does_not_mutate_the_original(tmp_path):
    template = bracket_template(tmp_path)
    heavy = template.replace(model_name="BRK-HEAVY", file_name="heavy.xlsx")
    assert heavy.model_name == "BRK-HEAVY"
    assert template.model_name == "BRK-MASTER"
    assert template.file_name == "brk_master_dt.xlsx"


def test_new_table_accepts_overrides(tmp_path):
    template = bracket_template(tmp_path)
    table = template.new_table(model_name="OTHER", columns=(LENGTH,))
    assert table.model_name == "OTHER"
    assert table.headers() == ("Length@Boss-Extrude1",)


def test_template_without_file_name_has_no_output_path():
    assert sw.TableTemplate(model_name="M").output_path is None


def test_blank_table_writes_a_valid_but_empty_file(tmp_path):
    """Headers, title and Family, no rows: the skeleton to fill in by hand."""
    destination = tmp_path / "skeleton.xlsx"
    sw.blank_table("BRK-MASTER", [LENGTH, WIDTH], path=destination)
    workbook = load_workbook(destination)
    sheet = workbook["Sheet1"]
    assert sheet["A1"].value == "Design Table for: BRK-MASTER"
    assert sheet["A2"].value is None
    assert sheet["B2"].value == "Length@Boss-Extrude1"
    assert sheet["A3"].value is None
    assert workbook.defined_names["Family"].attr_text == "Sheet1!$A$2"


def test_blank_table_can_seed_configuration_names(tmp_path):
    destination = tmp_path / "skeleton.xlsx"
    sw.blank_table("M", [LENGTH], path=destination, configurations=["A", "B"])
    sheet = load_workbook(destination)["Sheet1"]
    assert sheet["A3"].value == "A"
    assert sheet["A4"].value == "B"


def test_blank_table_returns_the_table_without_writing():
    table = sw.blank_table("M", [LENGTH])
    assert isinstance(table, sw.DesignTable)
    assert table.configurations == ()
