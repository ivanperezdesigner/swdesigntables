"""Turning a table model into an .xlsx file.

The only module that imports openpyxl. Everything above it is plain data, so
the model can be tested without Excel and a future reader or COM backend can
sit beside this one without disturbing anything.
"""

from __future__ import annotations

import re
from datetime import datetime
from typing import IO, TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from swdesigntables.values import MissingValue, normalize

if TYPE_CHECKING:
    from swdesigntables.table import DesignTable

__all__ = ["build_workbook", "FAMILY_NAME", "TITLE_PREFIX", "FIXED_TIMESTAMP"]

FAMILY_NAME = "Family"
TITLE_PREFIX = "Design Table for: "
FIRST_HEADER_COLUMN = 2
FIRST_DATA_ROW = 3

# A fixed stamp so two identical runs produce identical bytes. openpyxl would
# otherwise write datetime.now() into the document properties and every
# regenerated table would look modified in git.
FIXED_TIMESTAMP = datetime(2000, 1, 1, 0, 0, 0)

_SAFE_SHEET_REFERENCE = re.compile(r"^[A-Za-z_][A-Za-z0-9_.]*$")


def _sheet_reference(sheet_name: str) -> str:
    """Return ``Sheet1!$A$2``, quoting the sheet name when it needs it."""
    if _SAFE_SHEET_REFERENCE.match(sheet_name):
        return f"{sheet_name}!$A$2"
    escaped = sheet_name.replace("'", "''")
    return f"'{escaped}'!$A$2"


def build_workbook(table: DesignTable) -> Workbook:
    """Return the openpyxl workbook for ``table``, without writing anything."""
    workbook = Workbook()
    # A fresh workbook always has exactly one sheet; workbook.active is
    # typed as optional only because a loaded workbook might have none.
    sheet: Worksheet = workbook.worksheets[0]
    sheet.title = table.sheet_name

    sheet.cell(row=1, column=1, value=TITLE_PREFIX + table.model_name)

    headers = [col.header(table.vocabulary) for col in table.columns]
    for offset, header in enumerate(headers):
        sheet.cell(row=2, column=FIRST_HEADER_COLUMN + offset, value=header)

    for row_offset, config in enumerate(table.configurations):
        row = FIRST_DATA_ROW + row_offset
        sheet.cell(row=row, column=1, value=config.name)
        for col_offset, col in enumerate(table.columns):
            if col.key in config.values:
                value = config.values[col.key]
            elif table.missing is MissingValue.FILL:
                value = table.fill_value
            else:
                continue
            payload, force_text = normalize(
                value,
                state_format=table.state_format,
                round_floats=table.round_floats,
                value_kind=col.value_kind,
            )
            if payload is None:
                continue
            cell = sheet.cell(row=row, column=FIRST_HEADER_COLUMN + col_offset)
            cell.value = payload
            if force_text:
                # Without this openpyxl stores '=x' as a formula with no cached
                # result, and SOLIDWORKS reads an empty parameter.
                cell.data_type = "s"

    if table.autosize_columns:
        _autosize(sheet, headers, table)

    workbook.defined_names[FAMILY_NAME] = DefinedName(
        FAMILY_NAME, attr_text=_sheet_reference(table.sheet_name)
    )
    for name, ref in table.defined_names.items():
        workbook.defined_names[name] = DefinedName(name, attr_text=ref)

    for extra in table.extra_sheets:
        worksheet = workbook.create_sheet(title=extra.name)
        for row_index, row_values in enumerate(extra.rows, start=1):
            for col_index, value in enumerate(row_values, start=1):
                payload, force_text = normalize(
                    value,
                    state_format=table.state_format,
                    round_floats=table.round_floats,
                )
                if payload is None:
                    continue
                cell = worksheet.cell(row=row_index, column=col_index)
                cell.value = payload
                if force_text:
                    cell.data_type = "s"
        if extra.hidden:
            # 'hidden', never 'veryHidden': you must be able to unhide it in
            # Excel when the table misbehaves.
            worksheet.sheet_state = "hidden"

    workbook.properties.created = FIXED_TIMESTAMP
    workbook.properties.modified = FIXED_TIMESTAMP
    return workbook


def _autosize(sheet: Worksheet, headers: list[str], table: DesignTable) -> None:
    """Widen columns to fit their contents, for humans who open the file."""
    longest = max((len(config.name) for config in table.configurations), default=0)
    sheet.column_dimensions["A"].width = max(12, min(longest + 2, 60))
    for offset, header in enumerate(headers):
        letter = get_column_letter(FIRST_HEADER_COLUMN + offset)
        sheet.column_dimensions[letter].width = max(10, min(len(header) + 2, 60))


def save(table: DesignTable, target: str | IO[bytes]) -> None:
    """Write ``table`` to a path or a binary stream."""
    build_workbook(table).save(target)
